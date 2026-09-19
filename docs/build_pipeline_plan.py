"""Génère docs/pipeline-plan.pptx (livrable « plan du pipeline ») + PDF et PNG de contrôle via PowerPoint.

    python docs/build_pipeline_plan.py            # pptx seul
    python docs/build_pipeline_plan.py --export   # + PDF et PNG par slide (PowerPoint COM, Windows)

Les schémas sont dessinés en formes natives (lisibles et éditables dans PowerPoint).
"""

from __future__ import annotations

import subprocess
import sys
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DOCS = Path(__file__).resolve().parent
OUT = DOCS / "pipeline-plan.pptx"

# Palette « Midnight amber » : navy dominant, ambre pour les sorties, rouge pour la fraude
NAVY = RGBColor(0x14, 0x21, 0x3D)
NAVY_2 = RGBColor(0x1F, 0x33, 0x5C)
AMBER = RGBColor(0xFC, 0xA3, 0x11)
AMBER_DARK = RGBColor(0xB3, 0x6B, 0x00)
RED = RGBColor(0xE7, 0x4C, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x24, 0x2B)
MUTED = RGBColor(0x6B, 0x72, 0x80)
PALE = RGBColor(0xF3, 0xF4, 0xF8)
LINE = RGBColor(0xD9, 0xDD, 0xE7)
ICE = RGBColor(0xCA, 0xD3, 0xE6)

W, H = Inches(13.333), Inches(7.5)
HEAD = "Cambria"
BODY = "Calibri"
I = Inches


# ───────────────────────── helpers ─────────────────────────
def tb(slide, text, x, y, w, h, size=14, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, italic=False, margin=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if not margin:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def title(slide, text, dark=False):
    tb(slide, text, I(0.6), I(0.35), I(12.1), I(0.8), size=32, bold=True, color=WHITE if dark else NAVY, font=HEAD,
       margin=False)


def chip(slide, text, x, y, w, h, fill=NAVY_2, color=WHITE, size=11, bold=True):
    rect(slide, x, y, w, h, fill, radius=0.3)
    tb(slide, text, x, y, w, h, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
       margin=False)


def node(slide, text, x, y, w, h, fill=WHITE, line=NAVY, color=INK, size=11, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Boîte de diagramme : première ligne en gras (titre), suivantes en normal."""
    rect(slide, x, y, w, h, fill, line=line, shape=shape, radius=0.12)
    lines = text if isinstance(text, list) else [text]
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = I(0.08)
    tf.margin_top = tf.margin_bottom = I(0.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size + 1 if i == 0 else size)
        r.font.bold = i == 0
        r.font.color.rgb = color
        r.font.name = BODY
    return box


def seg(slide, x1, y1, x2, y2, color=MUTED, arrow=False, width=1.5):
    """Segment droit ; flèche à l'extrémité (x2, y2) si arrow."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def path(slide, points, color=MUTED, width=1.5):
    """Polyligne à coudes, flèche au dernier point."""
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        seg(slide, x1, y1, x2, y2, color, arrow=(x2, y2) == points[-1], width=width)


def label(slide, text, x, y, w=I(1.8), h=I(0.28), size=9.5, color=MUTED, fill=WHITE, align=PP_ALIGN.CENTER):
    """Étiquette d'arête sur fond opaque (masque la ligne dessous)."""
    rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE)
    tb(slide, text, x, y, w, h, size=size, color=color, italic=True, align=align, anchor=MSO_ANCHOR.MIDDLE,
       margin=False)


def group_box(slide, text, x, y, w, h, fill, line, color=NAVY):
    rect(slide, x, y, w, h, fill, line=line, radius=0.04)
    tb(slide, text, x + I(0.15), y + I(0.05), w - I(0.3), I(0.35), size=10.5, bold=True, color=color, margin=False)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ───────────────────────── slides ─────────────────────────
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    tb(s, "Automatic Fraud Detection", I(0.8), I(1.3), I(11.5), I(1.0), size=44, bold=True, color=WHITE, font=HEAD,
       margin=False)
    tb(s, "Plan du pipeline de détection de fraude en temps réel — ETL with Airflow", I(0.8), I(2.25), I(11.5), I(0.6),
       size=20, color=ICE, margin=False)
    tb(s, "Emeline ROBLOT · Jedha · septembre 2026 · github.com/emelineroblot/automatic-fraud-detection",
       I(0.8), I(2.85), I(11.5), I(0.4), size=13, color=RGBColor(0x9A, 0xA7, 0xC2), margin=False)

    stats = [("1 tx / min", "collectée depuis l'API et scorée par Airflow"),
             ("< 30 s", "entre la réception d'une transaction et l'alerte Discord"),
             ("0,805", "PR-AUC du RandomForest retenu — split temporel, 0,39 % de fraudes"),
             ("19", "ressources AWS déclarées en Terraform (EC2, S3, IAM, réseau, secrets)")]
    x0, y0, cw, ch, gap = I(0.8), I(4.2), I(2.75), I(1.9), I(0.25)
    for i, (big, small) in enumerate(stats):
        x = x0 + i * (cw + gap)
        rect(s, x, y0, cw, ch, NAVY_2)
        tb(s, big, x + I(0.2), y0 + I(0.2), cw - I(0.4), I(0.8), size=30, bold=True, color=AMBER, font=HEAD,
           margin=False)
        tb(s, small, x + I(0.2), y0 + I(1.0), cw - I(0.4), I(0.8), size=12, color=ICE, margin=False)
    notes(s, "Deux besoins métier : être notifié dès qu'une fraude est détectée, et disposer chaque matin du bilan "
             "des paiements et fraudes de la veille. Priorité de l'énoncé : le pipeline de données, pas l'algorithme.")
    return s


def slide_needs(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, "Deux besoins métier, deux pipelines Airflow")

    cols = [
        ("Besoin 1 — être notifié dès qu'une fraude est détectée",
         "DAG fraud_realtime · schedule * * * * * · run ≈ 18 s",
         ["extract", "transform", "predict", "load", "check_fraud", "alert | no_fraud"],
         ["extract : GET de l'API (JSON double-encodé), 2 retries",
          "transform : normalisation, current_time (ms) → trans_time UTC",
          "predict : modèle chargé depuis MLflow (models:/fraud_detection@production)",
          "load : INSERT … ON CONFLICT (trans_num) DO NOTHING RETURNING — dédoublonnage",
          "check_fraud : branche → alert (embed Discord + trace fraud_alerts) ou no_fraud",
          "Garde-fous : max_active_runs = 1, catchup = False, dagrun_timeout 5 min"],
         RED),
        ("Besoin 2 — chaque matin, les paiements et fraudes de la veille",
         "DAG fraud_daily_report · schedule 0 6 * * * (Europe/Paris)",
         ["aggregate", "export_csv", "archive_s3", "store", "notify"],
         ["aggregate : volumes, montants, top catégories / états, répartition horaire",
          "précision et rappel du jour vs label fourni par l'API (vérité terrain)",
          "export_csv : reports/transactions_<date>.csv → archive_s3 : s3://…/reports/",
          "store : table daily_reports (agrégats + payload JSONB)",
          "notify : embed Discord récapitulatif",
          "Paramètre report_date pour rejouer une journée (démo)"],
         AMBER),
    ]
    x0, y0, cw = I(0.6), I(1.4), I(5.95)
    for i, (head, sub, chips, bullets, accent) in enumerate(cols):
        x = x0 + i * (cw + I(0.25))
        rect(s, x, y0, cw, I(5.1), PALE)
        tb(s, head, x + I(0.3), y0 + I(0.2), cw - I(0.6), I(0.7), size=16, bold=True, color=NAVY, font=HEAD,
           margin=False)
        tb(s, sub, x + I(0.3), y0 + I(0.9), cw - I(0.6), I(0.35), size=11, color=MUTED, italic=True, margin=False)
        n = len(chips)
        gap = I(0.12)
        chip_w = (cw - I(0.6) - gap * (n - 1)) / n
        cy = y0 + I(1.35)
        for j, c in enumerate(chips):
            cx = x + I(0.3) + j * (chip_w + gap)
            chip(s, c, int(cx), cy, int(chip_w), I(0.42), fill=accent if j == n - 1 else NAVY_2, size=9.5)
        box = tb(s, bullets, x + I(0.3), y0 + I(2.0), cw - I(0.6), I(3.0), size=13, color=INK, margin=False)
        for p in box.text_frame.paragraphs:
            p.space_after = Pt(6)
    notes(s, "Le pipeline temps réel est un micro-batch à la minute : suffisant pour une API qui produit une "
             "transaction par minute. Le dédoublonnage par clé primaire rend le pipeline rejouable sans double alerte. "
             "Le rapport quotidien réutilise le label de l'API comme vérité terrain pour suivre la qualité du modèle.")
    return s


def slide_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, "Architecture logique")
    cx = [I(0.6), I(3.1), I(5.6), I(8.1), I(10.6)]        # 5 colonnes
    w, h = I(2.15), I(1.15)
    ya, yb, yc = I(1.55), I(3.35), I(5.15)                 # 3 rangées

    def mid(x):
        return x + w // 2

    def my(y):
        return y + h // 2

    group_box(s, "Docker Compose", cx[1] - I(0.2), ya - I(0.35), cx[4] + w - cx[1] + I(0.4), yc + h - ya + I(0.55),
              PALE, LINE)

    node(s, ["API temps réel", "HF Space", "1 transaction / minute"], cx[0], ya, w, h)
    node(s, ["fraudTest.csv", "555 719 transactions", "0,39 % de fraudes"], cx[0], yc, w, h, shape=MSO_SHAPE.CAN)
    node(s, ["Entraînement", "training/train.py", "split temporel · LR / RF / LightGBM"], cx[1], yc, w, h)
    node(s, ["MLflow 3", "tracking + registry", "alias production"], cx[2], yc, w, h)
    node(s, ["DAG fraud_realtime", "chaque minute", "extract → transform → predict", "→ load → branch"],
         cx[2], ya, w, h, fill=NAVY_2, line=NAVY_2, color=WHITE, size=10)
    node(s, ["PostgreSQL 16", "transactions · fraud_alerts", "daily_reports · dédup PK"], cx[3], ya, w, h,
         shape=MSO_SHAPE.CAN, size=10)
    node(s, ["DAG fraud_daily_report", "6h Europe/Paris", "aggregate → export_csv", "→ store → notify"],
         cx[3], yc, w, h, fill=NAVY_2, line=NAVY_2, color=WHITE, size=10)
    node(s, ["Dashboard Streamlit", ":8501"], cx[4], ya, w, h)
    node(s, ["Sorties", "Discord : alerte fraude,", "rapport quotidien", "CSV quotidien (S3 en prod)"],
         cx[4], yb, w, h, fill=AMBER, line=AMBER, color=NAVY)

    seg(s, cx[0] + w, my(ya), cx[2], my(ya), arrow=True)                                   # API -> RT
    label(s, "GET chaque minute", mid(cx[1]) - I(0.7), my(ya) - I(0.14), w=I(1.4))
    seg(s, cx[0] + w, my(yc), cx[1], my(yc), arrow=True)                                   # CSV -> TRAIN
    seg(s, cx[1] + w, my(yc), cx[2], my(yc), arrow=True)                                   # TRAIN -> MLF
    label(s, "log_model + alias", cx[1] + w - I(0.4), my(yc) + I(0.2), w=I(1.15))
    seg(s, mid(cx[2]), yc, mid(cx[2]), ya + h, arrow=True)                                 # MLF -> RT
    label(s, "load_model", mid(cx[2]) + I(0.1), my(yb) + I(0.3), w=I(0.9), align=PP_ALIGN.LEFT)
    seg(s, cx[2] + w, my(ya), cx[3], my(ya), arrow=True)                                   # RT -> PG
    seg(s, cx[3] + w, my(ya), cx[4], my(ya), arrow=True)                                   # PG -> DASH
    seg(s, mid(cx[3]), ya + h, mid(cx[3]), yc, arrow=True)                                 # PG -> DR
    label(s, "agrégats J-1", mid(cx[3]) + I(0.1), my(yb) + I(0.3), w=I(0.95), align=PP_ALIGN.LEFT)
    xr = mid(cx[2]) - I(0.6)
    path(s, [(xr, ya + h), (xr, my(yb)), (cx[4], my(yb))], color=RED)                     # RT -> OUT
    label(s, "alerte si is_fraud_pred = 1", xr + I(0.1), my(yb) - I(0.32), w=I(1.9), color=RED, align=PP_ALIGN.LEFT)
    xd = mid(cx[3]) + I(0.6)
    yl = yb + h + I(0.3)
    path(s, [(xd, yc), (xd, yl), (mid(cx[4]), yl), (mid(cx[4]), yb + h)], color=AMBER)     # DR -> OUT
    label(s, "rapport + CSV", xd + I(0.1), yl + I(0.05), w=I(1.05), color=AMBER_DARK, align=PP_ALIGN.LEFT)

    tb(s, "Un orchestrateur (Airflow), un entrepôt (PostgreSQL), un registre (MLflow) : le modèle est un artefact "
          "versionné ; le preprocessing (FeatureBuilder) est dans le pipeline sklearn loggé — une ligne brute de l'API "
          "traverse exactement les mêmes transformations qu'une ligne du CSV.",
       I(0.6), I(6.6), I(12.1), I(0.6), size=11.5, color=MUTED, italic=True, margin=False)
    notes(s, "Deux entrées : le dataset historique entre par l'entraînement dans MLflow ; l'API temps réel entre dans "
             "Airflow. Elles convergent dans la tâche predict, qui charge le modèle par alias. Tout ce qui sort "
             "(Discord, CSV, dashboard) part de PostgreSQL.")
    return s


def slide_aws(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    title(s, "Déploiement en production — AWS eu-north-1 (Terraform)")
    lw = I(2.3)
    node(s, ["Poste opérateur", "IP unique autorisée", "terraform apply · UIs 8080 / 5000 / 8501"],
         I(0.6), I(1.5), lw, I(1.0))
    node(s, ["API temps réel (HF Space)", "1 transaction / minute"], I(0.6), I(2.9), lw, I(0.8))
    node(s, ["GitHub", "automatic-fraud-detection", "cloné au boot (user_data)"], I(0.6), I(4.1), lw, I(0.9),
         shape=MSO_SHAPE.CAN)
    node(s, ["Discord", "alerte fraude · rapport"], I(0.6), I(5.5), lw, I(0.8), fill=AMBER, line=AMBER, color=NAVY)

    group_box(s, "AWS eu-north-1 · VPC par défaut · 19 ressources Terraform · security group = IP opérateur",
              I(3.4), I(1.45), I(9.35), I(5.55), PALE, LINE)
    group_box(s, "EC2 m7i-flex.large · Ubuntu 24.04 · EBS 30 Go chiffré · docker compose (docker/prod/)",
              I(3.7), I(1.95), I(6.1), I(4.8), WHITE, NAVY_2)
    node(s, ["Airflow 3", "api-server · scheduler · dag-processor", "DAG fraud_realtime", "DAG fraud_daily_report"],
         I(4.0), I(2.5), I(2.3), I(1.45), fill=NAVY_2, line=NAVY_2, color=WHITE, size=10)
    node(s, ["MLflow 3 server", "registry : fraud_detection@production"], I(7.15), I(2.5), I(2.4), I(0.95))
    node(s, ["training (profil compose)", "one-shot au boot · python 3.13"], I(7.15), I(4.05), I(2.4), I(0.65), size=10)
    node(s, ["PostgreSQL 16", "airflow · mlflow · fraud", "volume Docker sur l'EBS"], I(4.0), I(4.75), I(2.3), I(1.1),
         shape=MSO_SHAPE.CAN)
    node(s, ["Streamlit", ":8501"], I(7.15), I(5.1), I(2.4), I(0.6))
    node(s, ["S3", "chiffré · versionné", "data/fraudTest.csv", "mlflow-artifacts/ · reports/"],
         I(10.2), I(2.5), I(2.35), I(1.5), shape=MSO_SHAPE.CAN)
    node(s, ["IAM — rôle d'instance", "accès S3 uniquement", "aucune clé dans le code"], I(10.2), I(4.75), I(2.35),
         I(0.95))

    seg(s, I(2.9), I(2.0), I(3.4), I(2.0), arrow=True)                                        # OP -> AWS
    seg(s, I(2.9), I(3.3), I(4.0), I(3.3), arrow=True)                                        # API -> AF
    seg(s, I(2.9), I(4.55), I(3.7), I(4.55), arrow=True)                                      # GH -> EC2
    label(s, "git clone", I(2.95), I(4.62), w=I(0.7), size=8.5)
    path(s, [(I(4.0), I(3.75)), (I(3.2), I(3.75)), (I(3.2), I(5.9)), (I(2.9), I(5.9))], color=RED)  # AF -> Discord
    label(s, "webhook", I(3.27), I(5.5), w=I(0.7), size=8.5, color=RED)
    seg(s, I(7.15), I(2.95), I(6.3), I(2.95), arrow=True)                                     # MLF -> AF
    label(s, "load_model", I(6.35), I(2.62), w=I(0.75), size=8.5)
    seg(s, I(7.6), I(4.05), I(7.6), I(3.45), arrow=True)                                      # TRN -> MLF
    label(s, "log_model + alias", I(7.7), I(3.48), w=I(1.15), size=8.5, align=PP_ALIGN.LEFT)
    seg(s, I(5.15), I(3.95), I(5.15), I(4.75), arrow=True)                                    # AF -> PG
    label(s, "transactions · prédictions", I(5.25), I(4.22), w=I(1.5), size=8.5, align=PP_ALIGN.LEFT)
    seg(s, I(7.15), I(5.4), I(6.3), I(5.4), arrow=True)                                       # DASH -> PG
    seg(s, I(9.55), I(2.95), I(10.2), I(2.95), arrow=True)                                    # MLF -> S3
    label(s, "artefacts", I(9.57), I(2.62), w=I(0.6), size=8.5)
    path(s, [(I(10.9), I(4.0)), (I(10.9), I(4.37)), (I(9.55), I(4.37))])                      # S3 -> TRN
    label(s, "dataset", I(9.75), I(4.44), w=I(0.65), size=8.5)
    path(s, [(I(6.3), I(3.8)), (I(9.9), I(3.8)), (I(9.9), I(3.4)), (I(10.2), I(3.4))], color=AMBER)  # AF -> S3
    label(s, "CSV quotidien", I(8.85), I(3.85), w=I(1.0), size=8.5, color=AMBER_DARK)
    seg(s, I(11.9), I(4.75), I(11.9), I(4.02), arrow=True, color=LINE)                         # IAM -> S3

    tb(s, "RDS PostgreSQL était la cible : bloqué par le quota du plan gratuit du compte (1 instance, déjà utilisée). "
          "Coût ≈ 2,5 $ / jour · terraform destroy après la soutenance.",
       I(0.6), I(7.05), I(12.1), I(0.35), size=10.5, color=MUTED, italic=True, margin=False)
    notes(s, "Le même docker compose qu'en développement, sur une EC2 déclarée en Terraform. user_data : Docker → "
             "git clone → .env (secrets générés par Terraform) → compose up → entraînement → dépause des DAGs. "
             "Le dataset, les artefacts MLflow et les rapports sont dans S3 via le rôle d'instance.")
    return s


def slide_choices(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    title(s, "Choix d'architecture — et ce qui a été écarté", dark=True)
    rows = [
        ("Airflow seul, DAG à la minute", "Kafka + producer / consumer",
         "l'API produit 1 transaction par minute : Kafka ajouterait 2 services pour une latence identique. "
         "Pertinent si le débit monte ou si plusieurs consommateurs lisent le flux."),
        ("Modèle chargé depuis le registry MLflow", "API FastAPI de scoring",
         "un seul client (le DAG) ; l'alias production découple déjà entraînement et exploitation. "
         "Une API se justifie avec un second consommateur."),
        ("Preprocessing dans le pipeline sklearn", "code de features dupliqué dans le DAG",
         "un seul artefact, predict_proba sur du brut : aucune divergence training / inférence possible "
         "(exigence « réutilisable »)."),
        ("PostgreSQL", "warehouse analytique (BigQuery, Snowflake)",
         "1 440 lignes / jour : clé primaire pour le dédoublonnage, JSONB, agrégats SQL. "
         "Réplication vers un warehouse au-delà."),
        ("Webhook Discord", "e-mail SMTP",
         "5 lignes de code, message riche, visible en démo ; le métier « a juste besoin d'une notification »."),
        ("Terraform + docker compose sur EC2", "MWAA, ECS, RDS",
         "reproductible en une commande, ≈ 2,5 $ / jour ; RDS bloqué par le quota du compte, décrit comme cible."),
    ]
    x, y, w = I(0.6), I(1.35), I(12.1)
    hdr = [("Choix", I(3.0)), ("Alternative écartée", I(2.8)), ("Pourquoi", w - I(5.8))]
    cx = x
    for h, cw in hdr:
        tb(s, h, cx, y, cw, I(0.35), size=12, bold=True, color=AMBER, margin=False)
        cx += cw
    rh = I(0.9)
    for i, (a, b, c) in enumerate(rows):
        yy = y + I(0.45) + i * rh
        rect(s, x, yy, w, rh - I(0.08), NAVY_2)
        cx = x
        for txt, (h, cw), size, bold in zip((a, b, c), hdr, (12, 11.5, 10.5), (True, False, False)):
            tb(s, txt, cx + I(0.15), yy + I(0.06), cw - I(0.3), rh - I(0.2), size=size, bold=bold,
               color=WHITE if bold else ICE, anchor=MSO_ANCHOR.MIDDLE, margin=False)
            cx += cw
    notes(s, "Chaque choix est argumenté contre une alternative, avec le seuil à partir duquel l'alternative devient "
             "la bonne réponse. Limite assumée : l'API rejoue le dataset d'entraînement, la précision live est optimiste ; "
             "le label de l'API n'est jamais utilisé pour prédire.")
    return s


def build() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide_title(prs)
    slide_needs(prs)
    slide_architecture(prs)
    slide_aws(prs)
    slide_choices(prs)
    prs.save(OUT)
    print(f"{OUT.name}: {len(prs.slides)} slides")
    return OUT


# ───────────────────────── export PDF + PNG (PowerPoint COM) ─────────────────────────
def export(path: Path, png_dir: Path):
    import ctypes
    import ctypes.wintypes as wt

    import pythoncom
    import win32com.client as win32

    user32 = ctypes.windll.user32

    def dismiss(pid, seconds=3.0):
        deadline = time.time() + seconds
        while time.time() < deadline:
            found = []

            @ctypes.WINFUNCTYPE(ctypes.c_bool, wt.HWND, wt.LPARAM)
            def cb(h, _):
                p = wt.DWORD()
                user32.GetWindowThreadProcessId(h, ctypes.byref(p))
                cls = ctypes.create_unicode_buffer(64)
                user32.GetClassNameW(h, cls, 64)
                if p.value == pid and cls.value == "#32770" and user32.IsWindowVisible(h):
                    found.append(h)
                return True

            user32.EnumWindows(cb, 0)
            for h in found:
                user32.PostMessageW(h, 0x0010, 0, 0)
            time.sleep(0.5)

    pythoncom.CoInitialize()
    app = win32.DispatchEx("PowerPoint.Application")
    out = subprocess.run(["powershell", "-NoProfile", "-Command", "(Get-Process POWERPNT).Id"],
                         capture_output=True, text=True).stdout.split()
    pid = max(int(x) for x in out if x.isdigit())
    png_dir.mkdir(parents=True, exist_ok=True)
    try:
        pres = app.Presentations.Open(str(path), True, False, False)
        dismiss(pid)
        pres.SaveAs(str(path.with_suffix(".pdf")), 32)
        n = pres.Slides.Count
        for i in range(1, n + 1):
            pres.Slides(i).Export(str(png_dir / f"slide-{i}.png"), "PNG", 1920, 1080)
        pres.Close()
        print(f"PDF + {n} PNG -> {png_dir}")
    finally:
        app.Quit()


if __name__ == "__main__":
    out = build()
    if "--export" in sys.argv:
        idx = sys.argv.index("--export")
        png_dir = Path(sys.argv[idx + 1]) if len(sys.argv) > idx + 1 else DOCS / "_slides"
        export(out, png_dir)
